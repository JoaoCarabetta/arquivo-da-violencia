"""
Official violence data service (Ministry of Justice VDE data).

This module provides functions to:
1. Parse and ingest bancovde-YYYY.xlsx data (victim counts by municipality and month)
2. Resolve municipality names to 7-digit IBGE codes via ibge_population table
3. Calculate summed Formulário 1 totals (via formulario_1_indicators / map_natureza)
4. Query official statistics with window filtering

Data source: SINESP VDE (Validador de Dados Estatísticos)
URL pattern: https://www.gov.br/mj/pt-br/assuntos/sua-seguranca/seguranca-publica/estatistica/download/dnsp-base-de-dados/bancovde-YYYY.xlsx/@@download/file

File format: bancovde-2025.xlsx, sheet "2025", 14 columns
Headers: ["uf","municipio","evento","data_referencia","agente","arma","faixa_etaria","feminino","masculino","nao_informado","total_vitima","total","total_peso","abrangencia"]

Revision (preliminar vs consolidado):
The gov.br bancovde workbook does NOT encode revision in its columns. The portal
publishes one rolling snapshot per calendar year (bancovde-YYYY.xlsx). Operators
assign revision at ingest time via the ``revision`` kwarg (or CLI ``--revision``):
download the same file after preliminar consolidation → ingest with
``revision=preliminar``; re-download after homologation → ingest with
``revision=consolidado``. Both snapshots persist as separate rows keyed by the
5-tuple (code_muni, year_month, indicator, source_id, revision).
"""

from io import BytesIO
from typing import Dict, List, Any, BinaryIO, Union
from datetime import date, datetime, timedelta
from sqlmodel import select, func
from sqlmodel.ext.asyncio.session import AsyncSession
from loguru import logger

from app.models.official_violence_data import (
    OfficialRevision,
    OfficialSourceId,
    OfficialViolenceCount,
)
from app.services.official_typology import (
    formulario_1_indicators,
    map_natureza,
)

# gov.br portal (NOT broken dados.mj CKAN)
BANCOVDE_URL_TEMPLATE = (
    "https://www.gov.br/mj/pt-br/assuntos/sua-seguranca/seguranca-publica/"
    "estatistica/download/dnsp-base-de-dados/bancovde-{year}.xlsx/@@download/file"
)

# Expected bancovde-YYYY.xlsx headers (14 columns)
BANCOVDE_EXPECTED_HEADERS = [
    "uf", "municipio", "evento", "data_referencia", "agente", "arma",
    "faixa_etaria", "feminino", "masculino", "nao_informado",
    "total_vitima", "total", "total_peso", "abrangencia",
]

# v1 coverage window (issue #237 / #240)
BANCOVDE_WINDOW_START = "2025-09"


def bancovde_govbr_url(year: int) -> str:
    """Return the gov.br download URL for bancovde-YYYY.xlsx."""
    return BANCOVDE_URL_TEMPLATE.format(year=year)


def _excel_serial_to_year_month(serial_date: float) -> str:
    """
    Convert Excel serial date to YYYY-MM.

    Excel serial date: days since 1899-12-30 (Excel epoch)
    Examples: 45658 = 2025-01-01, 45689 = 2025-02-01

    Args:
        serial_date: Excel serial date number

    Returns:
        String in YYYY-MM format (e.g. "2025-09")
    """
    # Excel epoch is 1899-12-30
    excel_epoch = datetime(1899, 12, 30)
    parsed_date = excel_epoch + timedelta(days=int(serial_date))
    return parsed_date.strftime("%Y-%m")


def _data_referencia_to_year_month(value: Any) -> str | None:
    """
    Normalize bancovde ``data_referencia`` cell values to YYYY-MM.

    openpyxl with ``data_only=True`` may return Excel serial numbers (int/float)
    or native ``datetime``/``date`` objects depending on how the workbook was
    saved and evaluated.
    """
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.strftime("%Y-%m")

    if isinstance(value, date):
        return value.strftime("%Y-%m")

    try:
        return _excel_serial_to_year_month(float(value))
    except (ValueError, TypeError):
        return None


def parse_bancovde_workbook(
    workbook_bytes: Union[bytes, BinaryIO],
    year: int,
    since_year_month: str = BANCOVDE_WINDOW_START,
) -> List[Dict[str, Any]]:
    """
    Parse a bancovde-YYYY.xlsx workbook into row dicts.

    Uses the gov.br file layout: sheet named after the calendar year, 14 columns,
    Excel serial or datetime/date values in ``data_referencia``. Rows before
    ``since_year_month`` are
    dropped (default: 2025-09 inclusive window start).

    Args:
        workbook_bytes: Raw XLSX bytes or file-like object
        year: Calendar year (selects sheet name, e.g. 2025 → sheet "2025")
        since_year_month: Minimum YYYY-MM to keep (inclusive)

    Returns:
        List of row dicts keyed by column header names
    """
    from openpyxl import load_workbook

    if isinstance(workbook_bytes, bytes):
        source = BytesIO(workbook_bytes)
    else:
        source = workbook_bytes

    wb = load_workbook(source, read_only=True, data_only=True)
    sheet_name = str(year)

    if sheet_name not in wb.sheetnames:
        wb.close()
        raise ValueError(
            f"Sheet '{sheet_name}' not found in workbook. "
            f"Available sheets: {wb.sheetnames}"
        )

    ws = wb[sheet_name]

    headers: List[str] = []
    for col_idx in range(1, len(BANCOVDE_EXPECTED_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        header_value = cell.value
        if header_value is None:
            headers.append(f"col_{col_idx}")
        else:
            headers.append(str(header_value).strip())

    if headers != BANCOVDE_EXPECTED_HEADERS:
        logger.warning(
            "bancovde header mismatch for %s: expected %s, got %s",
            sheet_name,
            BANCOVDE_EXPECTED_HEADERS,
            headers,
        )

    records: List[Dict[str, Any]] = []
    skipped_count = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        record: Dict[str, Any] = {}
        for col_idx, value in enumerate(row):
            if col_idx >= len(headers):
                break
            record[headers[col_idx]] = value

        data_ref = record.get("data_referencia")
        year_month = _data_referencia_to_year_month(data_ref)
        if year_month is None:
            skipped_count += 1
            continue

        if year_month < since_year_month:
            skipped_count += 1
            continue

        records.append(record)

    wb.close()
    logger.info(
        "Parsed bancovde-%s.xlsx: kept %s rows >= %s (skipped %s)",
        year,
        len(records),
        since_year_month,
        skipped_count,
    )
    return records


async def download_bancovde_workbook(year: int) -> bytes:
    """
    Download bancovde-YYYY.xlsx from the gov.br portal.

    Network I/O only — use parse_bancovde_workbook() for parsing in tests.
    """
    import httpx

    url = bancovde_govbr_url(year)
    logger.info("Downloading bancovde-%s.xlsx from %s", year, url)
    async with httpx.AsyncClient(timeout=120.0) as client:
        response = await client.get(url)
        response.raise_for_status()
    logger.info("Downloaded %s bytes", len(response.content))
    return response.content


async def ingest_official_violence_data(
    session: AsyncSession,
    vde_data: List[Dict[str, Any]],
    source: str = "SINESP VDE",
    source_id: OfficialSourceId = OfficialSourceId.VALIDADOR,
    revision: OfficialRevision = OfficialRevision.CONSOLIDADO,
) -> None:
    """
    Ingest bancovde-YYYY.xlsx data (victim counts by municipality and month).

    Expected bancovde format (14 columns):
    - uf: State abbreviation (e.g. "SP")
    - municipio: Municipality name (e.g. "SÃO PAULO")
    - evento: Crime type (must match INDICATOR_MAPPING keys)
    - data_referencia: Excel serial date for first of month (e.g. 45901 = 2025-09-01)
    - agente, arma, faixa_etaria: Disaggregation dimensions (ignored for our totals)
    - feminino, masculino, nao_informado: Sex-disaggregated counts (unused - we use total_vitima)
    - total_vitima: Total victim count (sex-combined) - THIS IS WHAT WE USE
    - total, total_peso, abrangencia: Other columns (unused)

    Municipality resolution:
    - No 7-digit IBGE code in the file
    - Resolve (uf, municipio) → code_muni via ibge_population table
    - Drop rows that don't match an IBGE code

    Aggregation:
    - Rows are sliced by agente/arma/faixa_etaria
    - SUM total_vitima for same (uf, municipio, year_month, evento) before storing

    This function:
    1. Resolves municipality names to IBGE codes
    2. Groups and sums by (code_muni, year_month, indicator)
    3. Stores per-indicator counts (4 Formulário 1 types + intervenção separately)
    4. Is idempotent: re-ingesting the same month updates existing rows
    
    Note: Does not store a convenience total. Official municipal total is computed
    at query time by summing the four Formulário 1 types.

    Args:
        session: Database session
        vde_data: List of bancovde data rows (dicts with 14 columns)
        source: Data source description
        source_id: Official data source identity (default: validador)
        revision: Data revision stage (default: consolidado)
    """
    if not vde_data:
        return

    # Import here to avoid circular dependency
    from app.models.ibge_population import IBGEPopulation

    # First pass: collect all unique (uf, municipio) pairs for batch lookup
    unique_municipalities = set()
    for row in vde_data:
        uf = row.get("uf", "").strip()
        municipio = row.get("municipio", "").strip()
        if uf and municipio:
            unique_municipalities.add((municipio, uf))

    # Batch lookup IBGE codes with case-insensitive matching
    # Dump uses "SÃO PAULO", IBGE has "São Paulo" - need casefold comparison
    # Do NOT change lookup_city_codes globally (other callers rely on exact match)
    query = select(IBGEPopulation.code_muni, IBGEPopulation.name_muni, IBGEPopulation.abbrev_state)
    result = await session.execute(query)
    ibge_records = result.all()

    # Build case-insensitive lookup: (name_casefold, state_casefold) → code_muni
    code_lookup: Dict[tuple, int] = {}
    for code_muni, name_muni, abbrev_state in ibge_records:
        if name_muni and abbrev_state:
            key = (name_muni.strip().casefold(), abbrev_state.strip().casefold())
            code_lookup[key] = code_muni

    # Second pass: group and sum by (code_muni, year_month, indicator)
    grouped: Dict[tuple, int] = {}

    for row in vde_data:
        uf = row.get("uf", "").strip()
        municipio = row.get("municipio", "").strip()

        if not uf or not municipio:
            logger.debug(f"Missing uf/municipio in row, skipping")
            continue

        # Resolve to IBGE code (case-insensitive)
        lookup_key = (municipio.casefold(), uf.casefold())
        code_muni = code_lookup.get(lookup_key)
        if not code_muni:
            logger.debug(f"Could not resolve IBGE code for {municipio}/{uf}, skipping")
            continue

        # Parse evento via shared typology map
        evento = row.get("evento", "").strip()
        mapped = map_natureza(evento)
        if mapped["kind"] == "unmapped":
            continue

        indicator = mapped["indicator"]

        # Parse date
        data_ref = row.get("data_referencia")
        year_month = _data_referencia_to_year_month(data_ref)
        if year_month is None:
            logger.debug(f"Invalid or missing data_referencia in row: {data_ref!r}, skipping")
            continue

        # Parse victim count
        total_vitima = row.get("total_vitima")
        if total_vitima is None or total_vitima == "":
            victim_count = 0
        else:
            try:
                victim_count = int(float(total_vitima))
            except (ValueError, TypeError):
                logger.warning(f"Invalid total_vitima value: {total_vitima}")
                continue

        # Group key
        key = (code_muni, year_month, indicator)
        grouped[key] = grouped.get(key, 0) + victim_count

    n_muni_months = await upsert_official_indicator_counts(
        session,
        grouped,
        source_id=source_id,
        revision=revision,
        source=source,
    )
    logger.info(f"Ingested official violence data for {n_muni_months} municipality-months")


async def upsert_official_indicator_counts(
    session: AsyncSession,
    grouped: Dict[tuple, int],
    *,
    source_id: OfficialSourceId,
    revision: OfficialRevision,
    source: str,
) -> int:
    """
    Upsert official counts keyed by (code_muni, year_month, indicator).

    Unique store key is the 5-tuple including source_id and revision (#238).
    Returns the number of distinct municipality-months written.
    """
    municipality_months: set[tuple] = set()

    for (code_muni, year_month, indicator), victim_count in grouped.items():
        municipality_months.add((code_muni, year_month))
        query_existing = select(OfficialViolenceCount).where(
            OfficialViolenceCount.code_muni == code_muni,
            OfficialViolenceCount.year_month == year_month,
            OfficialViolenceCount.indicator == indicator,
            OfficialViolenceCount.source_id == source_id,
            OfficialViolenceCount.revision == revision,
        )
        result = await session.execute(query_existing)
        existing = result.scalar_one_or_none()

        if existing:
            existing.victim_count = victim_count
            existing.updated_at = datetime.utcnow()
        else:
            session.add(
                OfficialViolenceCount(
                    code_muni=code_muni,
                    year_month=year_month,
                    indicator=indicator,
                    source_id=source_id,
                    revision=revision,
                    victim_count=victim_count,
                    is_total=False,
                    source=source,
                )
            )

    await session.commit()
    return len(municipality_months)


async def get_official_violence_totals(
    session: AsyncSession,
    code_munis: List[int],
    min_year_month: str = BANCOVDE_WINDOW_START,
    source_id: OfficialSourceId = OfficialSourceId.VALIDADOR,
    revision: OfficialRevision = OfficialRevision.CONSOLIDADO,
) -> List[Dict[str, Any]]:
    """
    Get official municipal totals (Formulário 1 types only) for municipalities.
    
    Sums the four exclusive Formulário 1 types at query time:
    - homicidio_doloso
    - feminicidio
    - latrocinio
    - lesao_corporal_seguida_morte

    Args:
        session: Database session
        code_munis: List of IBGE municipal codes
        min_year_month: Minimum year-month (YYYY-MM) to include (default: 2025-09)
        source_id: Official data source identity (default: validador)
        revision: Data revision stage (default: consolidado)

    Returns:
        List of dicts with keys: code_muni, year_month, victim_count
    """
    if not code_munis:
        return []
    
    formulario_1_types = formulario_1_indicators()

    query = select(
        OfficialViolenceCount.code_muni,
        OfficialViolenceCount.year_month,
        func.sum(OfficialViolenceCount.victim_count).label("victim_count")
    ).where(
        OfficialViolenceCount.code_muni.in_(code_munis),
        OfficialViolenceCount.indicator.in_(formulario_1_types),
        OfficialViolenceCount.year_month >= min_year_month,
        OfficialViolenceCount.source_id == source_id,
        OfficialViolenceCount.revision == revision,
    ).group_by(
        OfficialViolenceCount.code_muni,
        OfficialViolenceCount.year_month
    ).order_by(
        OfficialViolenceCount.code_muni,
        OfficialViolenceCount.year_month
    )

    result = await session.execute(query)
    rows = result.all()

    return [
        {
            "code_muni": r.code_muni,
            "year_month": r.year_month,
            "victim_count": r.victim_count
        }
        for r in rows
    ]
